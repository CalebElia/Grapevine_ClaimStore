"""Build the reference the benchmark scores against.

EVERY ELEMENT CARRIES ITS EVIDENCE TIER, because they are not equally trustworthy and a
score is only as good as what produced it:

  verified   a human read/heard it and wrote the exact words   — Garber's 476 words,
             Stults' 1,270. This is real gold.
  confirmed  a human affirmed a machine's proposal             — the 20 cluster names.
             Gold for WHO, not for WHERE the boundaries fall.
  reference  machine output, human-corrected but not audited   — the turn boundaries
             and timings from S1-S3. NOT gold. The incumbent produced these, so any
             metric scored against them flatters the incumbent and must say so.

WHY THE WINDOW IS 0:00-8:00. It contains the two hardest and most diagnostic things in
the meeting:
  0:30-2:25  ROLL CALL — 13 speaker changes in under two minutes, nearly all utterances
             under 2 seconds. This is where our pipeline produced 38% sub-1.5s segments,
             where 67% of interjections carried nothing, and where a junk cluster got
             confidently named. If a vendor handles roll call, it handles this corpus.
  4:11-7:20  KEN GARBER's public comment — three minutes of continuous single-speaker
             speech whose exact words are known, containing 'A2Zero' twice (which Gemini
             rendered 'A20' both times) and the date 2050 (which we rendered 2015).

Regenerate:  python -m pipeline.bench.ground_truth --video lWvRVUMyLP4
Then EDIT the emitted file by hand to upgrade tiers as you verify more.
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

REPO = Path(__file__).parent.parent.parent
WINDOW = (0.0, 480.0)

# Terms whose failure is silent and expensive, with the misrenderings actually observed.
ENTITIES = {
    "A2Zero":            ["A20", "A2 zero", "A to zero", "a20"],
    "Missy Stults":      ["Stoltz", "Staltz", "Stultz"],
    "AmeriCorps":        ["Americanore", "America Corps"],
    "DTE":               ["DTE Energy Company"],
    "Bryant":            ["Briant", "Brian"],
    "Solarize":          ["Solar Eyes", "Solarized"],
    "ARCA":              ["Arca", "Arcka"],
    "Scio Township":     ["Scyo Township", "Sile Township"],
    "Ken Garber":        ["Ken Garver", "Ken Gerber"],
    "Colvin-Garcia":     ["Colvin Garcia", "Calvin Garcia"],
}


def build(video_id: str = "lWvRVUMyLP4", window=WINDOW) -> dict:
    from openpyxl import load_workbook
    import pipeline.export_review as ex
    C = ex.COL
    d = REPO / "processing" / video_id
    turns = json.loads((d / "turns.json").read_text())["turns"]
    wb = load_workbook(next((d / "review").glob("*.xlsx")), data_only=True)
    ws, sp = wb["Transcript"], wb["Speakers"]

    # confirmed cluster -> name, and clusters the human REJECTED as not-a-person
    names, rejected = {}, []
    for r in sp.iter_rows(min_row=2):
        if not r[0].value:
            continue
        cl, conf, corr, auto = r[0].value, str(r[8].value or "").lower(), r[9].value, r[1].value
        if conf == "n":
            rejected.append(cl)
        elif conf == "y" and (corr or auto):
            names[cl] = corr or auto

    # per-turn speaker overrides and verified text, keyed by utterance_key
    overrides, verified = {}, {}
    for r in ws.iter_rows(min_row=2):
        k = r[C["utterance_key"] - 1].value
        if not k:
            continue
        if r[C["✎ Speaker correction"] - 1].value:
            overrides[str(k)] = r[C["✎ Speaker correction"] - 1].value
        fix = r[C["✎ Text correction"] - 1].value
        if fix and len(str(fix).split()) > 50:      # substantial hand-verified passages
            verified[str(k)] = str(fix)

    lo, hi = window
    # THE REFERENCE TIMELINE MUST NOT OVERLAP ITSELF. An interjection is stored inside
    # its parent turn's span, so emitting both verbatim makes the parent look like it
    # contains two speakers everywhere an interjection occurs — which inflates every
    # vendor's welded_rate and depresses purity, including vendors that did nothing
    # wrong. DER assumes an exclusive timeline anyway. So interjections CARVE OUT their
    # span from the parent.
    spans, cuts = [], []
    for t in turns:
        if t["end"] < lo or t["start"] > hi:
            continue
        who = overrides.get(t["key"]) or names.get(t["speaker"])
        if who:
            spans.append([t["start"], t["end"], who, t["speaker"]])
        for j in t["interjections"]:
            if not (lo <= j["start"] <= hi) or not j["text"].strip():
                continue
            ik = f"{t['key']}#int@{int(j['start']*1000)}"
            jwho = overrides.get(ik) or names.get(j["speaker"])
            if jwho:
                cuts.append([j["start"], j["end"], jwho, j["speaker"]])

    exclusive = []
    for s, e, who, cl in spans:
        pieces = [(s, e)]
        for cs, ce, _, _ in cuts:
            nxt = []
            for ps, pe in pieces:
                if ce <= ps or cs >= pe:
                    nxt.append((ps, pe))
                else:
                    if ps < cs:
                        nxt.append((ps, cs))
                    if ce < pe:
                        nxt.append((ce, pe))
            pieces = nxt
        for ps, pe in pieces:
            if pe - ps > 0.15:            # drop slivers left by the carve
                exclusive.append((ps, pe, who, cl))
    exclusive.extend((s, e, who, cl) for s, e, who, cl in cuts)

    # STILL NOT EXCLUSIVE. Carving interjections out of their own parent is not enough —
    # S3 turns overlap EACH OTHER too (pyannote emits overlapping speech, and adjacent
    # turns share boundary instants). The first version of this left the reference
    # covering 661s of a 480s window, i.e. 138%, which made DER meaningless because
    # "who is speaking at time t" had more than one answer.
    #
    # Flatten onto atomic intervals and let the MOST SPECIFIC segment win — the shortest
    # one covering that instant. A brief interjection inside a long turn is the more
    # precise statement about who is speaking, so it should take the instant.
    edges = sorted({round(v, 3) for s, e, _, _ in exclusive for v in (s, e)})
    flat = []
    for a, b in zip(edges, edges[1:]):
        if b - a < 0.02:
            continue
        mid = (a + b) / 2
        covering = [(e - s, who, cl) for s, e, who, cl in exclusive if s <= mid < e]
        if not covering:
            continue
        _, who, cl = min(covering)
        if flat and flat[-1][2] == who and abs(flat[-1][1] - a) < 0.05:
            flat[-1][1] = b                      # extend the run
        else:
            flat.append([a, b, who, cl])
    # Clip to the window. A turn that starts inside it and runs long — Missy Stults'
    # presentation begins at 473s and continues to 756s — otherwise contributes its whole
    # duration to the reference, which is how coverage read 138% of a 480s window.
    segments = [{"start": round(max(s, lo), 2), "end": round(min(e, hi), 2),
                 "speaker": who, "tier": "confirmed", "cluster": cl}
                for s, e, who, cl in flat if min(e, hi) - max(s, lo) > 0.02]

    # verified passages that fall inside the window
    passages = []
    for t in turns:
        if t["key"] in verified and lo <= t["start"] <= hi:
            passages.append({"start": round(t["start"], 2), "end": round(t["end"], 2),
                             "cluster": t["speaker"],
                             "speaker": names.get(t["speaker"]),
                             "text": verified[t["key"]], "tier": "verified"})

    # timestamp reference: our own forced alignment. NOT GOLD — see module docstring.
    tr = json.loads((d / "transcript.json").read_text())["segments"]
    ref_words = [[w["word"], round(w["start"], 3)]
                 for s in tr for w in s.get("words", [])
                 if "start" in w and lo <= w["start"] <= hi]

    return {
        "video_id": video_id,
        "window": {"start": lo, "end": hi,
                   "why": "roll call (hardest diarization) + a fully verified passage"},
        "speaker_segments": segments,
        "speaker_names": sorted({s["speaker"] for s in segments}),
        "rejected_clusters": rejected,
        "verified_passages": passages,
        "entities": ENTITIES,
        "timestamp_reference_words": ref_words,
        "_tiers": {
            "verified": "human wrote the exact words — gold",
            "confirmed": "human affirmed a machine proposal — gold for WHO, not WHERE",
            "reference": "machine output, human-corrected, NOT audited — flatters the "
                         "incumbent on timing metrics",
        },
        "_caveats": [
            "speaker_segments boundaries come from pyannote+S3 and are tier 'reference' "
            "even though the LABELS are confirmed. DER is therefore approximate.",
            "timestamp_reference_words is WhisperX forced alignment. WhisperX is also a "
            "candidate. Do not let timestamp accuracy decide the bake-off.",
        ],
    }


def load(video_id: str = "lWvRVUMyLP4") -> dict:
    p = REPO / "processing" / video_id / "bench" / "ground_truth.json"
    if not p.exists():
        raise FileNotFoundError(f"{p} — run: python -m pipeline.bench.ground_truth")
    return json.loads(p.read_text())


def main() -> int:
    ap = argparse.ArgumentParser(description="build the benchmark reference")
    ap.add_argument("--video", default="lWvRVUMyLP4")
    ap.add_argument("--start", type=float, default=WINDOW[0])
    ap.add_argument("--end", type=float, default=WINDOW[1])
    a = ap.parse_args()
    gt = build(a.video, (a.start, a.end))
    out = REPO / "processing" / a.video / "bench" / "ground_truth.json"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(gt, indent=2))
    print(f"[bench] {out}")
    print(f"[bench] window {a.start:.0f}-{a.end:.0f}s | "
          f"{len(gt['speaker_segments'])} speaker segments, "
          f"{len(gt['speaker_names'])} distinct people, "
          f"{len(gt['verified_passages'])} verified passage(s), "
          f"{len(gt['timestamp_reference_words'])} reference word timings")
    for c in gt["_caveats"]:
        print(f"[bench] CAVEAT {c}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
