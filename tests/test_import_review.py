"""S4c — the review round-trip.

The properties under test are the ones whose failure is SILENT. A dropped correction,
a correction attached to the wrong turn, or a voiceprint filed under the wrong person
all produce a clean run and a plausible artifact. Nothing here checks that the happy
path works; it checks that the failure modes are loud.

The DB-backed tests build a throwaway database and skip when no server is reachable,
matching test_schema_runtime.py.
"""
from __future__ import annotations

import json
import os
import uuid
from pathlib import Path

import pytest

from pipeline.import_review import (
    TRANSCRIPT_HEADER, SPEAKERS_HEADER, name_key, norm_name,
    parse_workbook, apply, reanchor, resolve_people,
)

openpyxl = pytest.importorskip("openpyxl")
REPO = Path(__file__).parent.parent
SCHEMA_DIR = REPO / "schema"
HOST = os.environ.get("GRAPEVINE_PGHOST", "/tmp")
PORT = os.environ.get("GRAPEVINE_PGPORT", "5433")
USER = os.environ.get("GRAPEVINE_PGUSER", "grapevine")


# ───────────────────────── fixtures ─────────────────────────

TURNS = [
    {"seq": 0, "speaker": "SPEAKER_00", "start": 10.0, "end": 20.0,
     "text": "Good evening and welcome.", "key": "v:00000:10000-20000", "interjections": []},
    {"seq": 1, "speaker": "SPEAKER_01", "start": 21.0, "end": 30.0,
     "text": "Thank you Chair.", "key": "v:00001:21000-30000", "interjections": []},
    {"seq": 2, "speaker": "SPEAKER_00", "start": 31.0, "end": 45.0,
     "text": "The A20 goal is important.", "key": "v:00002:31000-45000", "interjections": []},
]
SPEAKERS = {"SPEAKER_00": {"name": "Brooks A Curtis", "method": "roll_call",
                           "evidence": "called", "confidence": 0.85}}


def _workbook(tmp_path: Path, transcript_rows, speaker_rows,
              t_header=None, s_header=None) -> Path:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Transcript"
    ws.append(t_header or TRANSCRIPT_HEADER)
    for r in transcript_rows:
        ws.append(r)
    sp = wb.create_sheet("Speakers")
    sp.append(s_header or SPEAKERS_HEADER)
    for r in speaker_rows:
        sp.append(r)
    p = tmp_path / "annotated.xlsx"
    wb.save(p)
    return p


def _trow(t, speaker=None, text=None, flag=None, note=None, key=..., name="Brooks A Curtis"):
    # Cluster and name are SEPARATE columns — see TRANSCRIPT_COLS.
    return [t["seq"], "0:00:10", t["speaker"], name, "roll call", 0.85, t["text"],
            speaker, text, flag, note, t["key"] if key is ... else key]


def _irow(parent, j, speaker=None, note=None):
    """An interjection row, with the synthetic '<turnkey>#int@<ms>' key."""
    return [None, "0:00:12", j["speaker"], "—", "interjection", None, "↳ " + j["text"],
            speaker, None, None, note, f"{parent['key']}#int@{int(j['start'] * 1000)}"]


def _srow(cluster, auto, how, confirm=None, correct=None, note=None):
    return [cluster, auto, how, 0.85, 5.0, 3, "0:00:10", "called", confirm, correct, note]


# ───────────────────────── name matching ─────────────────────────

def test_name_key_ignores_middle_initial():
    """The roster says 'Brooks A Curtis'; persons says 'Brooks Curtis'. Same person."""
    assert name_key("Brooks A Curtis") == name_key("Brooks Curtis")
    assert name_key("Sara M Nedrich") == name_key("Sara Nedrich")


def test_name_key_does_not_collapse_different_first_names():
    """The whole point of keying on (first, last): a shared surname is not a match."""
    assert name_key("Ken Garber") != name_key("Judah Garber")
    assert name_key("Missy Stults") != name_key("Melissa Stults")


def test_norm_name_strips_punctuation_and_case():
    assert norm_name("  Colvin-Garcia,  Carlene ") == "colvin garcia carlene"


# ───────────────────────── parsing ─────────────────────────

def test_reordered_header_is_a_hard_stop(tmp_path):
    """Reading by position after a column moved would attribute corrections to the
    wrong field — worse than failing, because it looks like it worked."""
    bad = list(TRANSCRIPT_HEADER)
    bad[6], bad[7] = bad[7], bad[6]
    wb = _workbook(tmp_path, [_trow(TURNS[0])], [], t_header=bad)
    with pytest.raises(ValueError, match="header does not match"):
        parse_workbook(wb, TURNS)


def test_confirm_semantics(tmp_path):
    rows = [
        _srow("SPEAKER_00", "Brooks A Curtis", "roll call", confirm="y"),
        _srow("SPEAKER_01", "Wrong Person", "addressed by chair", confirm="n"),
        _srow("SPEAKER_02", "Someone", "roll call"),                       # blank
        _srow("SPEAKER_03", None, "UNRESOLVED", confirm="y", correct="Anya Dale"),
    ]
    rep = parse_workbook(_workbook(tmp_path, [], rows), TURNS)
    d = {c.cluster: c for c in rep.clusters}
    assert d["SPEAKER_00"].decision == "confirmed" and d["SPEAKER_00"].name == "Brooks A Curtis"
    assert d["SPEAKER_01"].decision == "rejected" and d["SPEAKER_01"].name is None
    assert d["SPEAKER_02"].decision == "unreviewed"
    assert d["SPEAKER_03"].decision == "confirmed" and d["SPEAKER_03"].name == "Anya Dale"


def test_blank_confirm_is_never_agreement(tmp_path):
    """A name typed without confirming is applied to the transcript but must not
    become eligible for enrollment. Absence of an act is not consent."""
    rows = [_srow("SPEAKER_00", "Brooks A Curtis", "roll call", correct="Margaret Halpern")]
    rep = parse_workbook(_workbook(tmp_path, [], rows), TURNS)
    c = rep.clusters[0]
    assert c.name == "Margaret Halpern"
    assert c.decision == "unreviewed"
    assert any("Confirm is blank" in n for n in rep.notes)


def test_confirmed_with_no_name_anywhere_is_downgraded(tmp_path):
    rows = [_srow("SPEAKER_09", None, "UNRESOLVED", confirm="y")]
    rep = parse_workbook(_workbook(tmp_path, [], rows), TURNS)
    assert rep.clusters[0].decision == "unreviewed"


def test_unknown_flag_is_kept_not_dropped(tmp_path):
    """Rule 1 posture: discarding review work to enforce a closed list is the thing
    the vocabulary trigger exists to avoid."""
    wb = _workbook(tmp_path, [_trow(TURNS[0], flag="crosstalk")], [])
    rep = parse_workbook(wb, TURNS)
    assert rep.turns[0].flag == "crosstalk"
    assert rep.unknown_flags[0]["flag"] == "crosstalk"


def test_orphan_keys_are_reported_not_silently_skipped(tmp_path):
    rows = [_trow(TURNS[0], note="lost", key="v:99999:1-2"),
            _trow(TURNS[1], note="no key", key=None)]
    rep = parse_workbook(_workbook(tmp_path, rows, []), TURNS)
    assert len(rep.orphan_keys) == 2
    assert rep.turns == []
    assert {o["reason"] for o in rep.orphan_keys} == {
        "no turn in turns.json with this key", "utterance_key cell is empty"}


def test_edit_to_machine_column_is_detected(tmp_path):
    from pipeline.import_review import COL_TEXT
    row = _trow(TURNS[0], flag="check")
    row[COL_TEXT] = "I typed over the machine text"
    rep = parse_workbook(_workbook(tmp_path, [row], []), TURNS)
    assert rep.tampered[0]["column"] == "Transcript text"


def test_cluster_and_name_are_separate_columns():
    """Collapsing them hides the cluster for every NAMED speaker, which destroys the
    only signal that catches a bad identification: two turns a reviewer knows are
    different people sharing one cluster id. This is how SPEAKER_11 — a residual bin of
    sub-second fragments labelled 'Sara M Nedrich' — stayed invisible through a full
    roll-call review."""
    assert TRANSCRIPT_HEADER.index("Cluster") < TRANSCRIPT_HEADER.index("Speaker (auto)")
    assert "Cluster" in TRANSCRIPT_HEADER and "Speaker (auto)" in TRANSCRIPT_HEADER


def test_interjection_anchors_to_itself_not_its_parent(tmp_path):
    """An interjection's anchor must describe the interjection. Anchoring it to the
    containing turn makes it re-anchor onto that turn on any re-export, which put a
    roll-call answer ('Mike Berkowitz') into the clerk's turn."""
    parent = {"seq": 0, "speaker": "SPEAKER_04", "start": 36.0, "end": 42.0,
              "text": "Commissioner Berkowitz?", "key": "v:00000:36000-42000",
              "interjections": [{"speaker": "SPEAKER_07", "start": 38.8, "end": 40.0,
                                 "text": "Present in Ann Arbor."}]}
    j = parent["interjections"][0]
    rep = parse_workbook(
        _workbook(tmp_path, [_irow(parent, j, speaker="Mike Berkowitz")], []), [parent])
    a = rep.turns[0].anchor
    assert a["is_interjection"] is True
    assert a["cluster"] == "SPEAKER_07"                 # not SPEAKER_04
    assert (a["start_ms"], a["end_ms"]) == (38800, 40000)   # not 36000-42000


def test_interjection_with_no_matching_offset_is_an_orphan(tmp_path):
    parent = {"seq": 0, "speaker": "SPEAKER_04", "start": 36.0, "end": 42.0,
              "text": "x", "key": "v:00000:36000-42000", "interjections": []}
    row = _irow(parent, {"speaker": "SPEAKER_07", "start": 38.8, "end": 40.0, "text": "y"},
                speaker="Mike Berkowitz")
    rep = parse_workbook(_workbook(tmp_path, [row], []), [parent])
    assert rep.turns == []
    assert "no interjection at that offset" in rep.orphan_keys[0]["reason"]


def test_repeated_per_turn_rename_suggests_the_speakers_sheet(tmp_path):
    many = [{**TURNS[0], "seq": i, "key": f"v:{i:05d}:{i}000-{i+1}000",
             "start": float(i), "end": float(i) + 1} for i in range(6)]
    rows = [_trow(t, speaker="Missy Stults") for t in many]
    rep = parse_workbook(_workbook(tmp_path, rows, []), many)
    assert any("Speakers sheet" in n for n in rep.notes)


# ───────────────────────── applying ─────────────────────────

def test_turns_json_is_never_modified(tmp_path):
    """The contract that lets S3 be re-run without destroying human review."""
    d = tmp_path / "proc"
    d.mkdir()
    src = json.dumps({"video_id": "v", "turns": TURNS})
    (d / "turns.json").write_text(src)
    wb = _workbook(tmp_path, [_trow(TURNS[0], text="Good evening, welcome.")],
                   [_srow("SPEAKER_00", "Brooks A Curtis", "roll call", confirm="y")])
    rep = parse_workbook(wb, TURNS)
    apply("v", rep, TURNS, SPEAKERS, wb, d)
    assert (d / "turns.json").read_text() == src
    assert (d / "corrections.json").exists()
    assert (d / "turns.reviewed.json").exists()


def test_speakers_json_is_never_modified(tmp_path):
    """S4's output is this stage's INPUT. An early version of apply() overwrote it and
    destroyed machine-derived identifications that the workbook could not reconstruct
    (a chair_address match at 0.70, and a cluster that S4 left unresolved). A stage
    that clobbers its own input turns a bad import from wrong into unrecoverable."""
    d = tmp_path / "proc"; d.mkdir()
    src = json.dumps(SPEAKERS)
    (d / "speakers.json").write_text(src)
    wb = _workbook(tmp_path, [], [_srow("SPEAKER_00", "Brooks A Curtis", "roll call",
                                        confirm="y", correct="Somebody Else")])
    rep = parse_workbook(wb, TURNS)
    apply("v", rep, TURNS, SPEAKERS, wb, d)
    assert (d / "speakers.json").read_text() == src
    assert json.loads((d / "speakers.reviewed.json").read_text())["SPEAKER_00"]["name"] \
        == "Somebody Else"


def test_text_correction_marks_alignment_stale(tmp_path):
    d = tmp_path / "proc"; d.mkdir()
    wb = _workbook(tmp_path, [_trow(TURNS[2], text="The A2Zero goal is important.")], [])
    rep = parse_workbook(wb, TURNS)
    apply("v", rep, TURNS, SPEAKERS, wb, d)
    t = json.loads((d / "turns.reviewed.json").read_text())["turns"][2]
    assert t["text"] == "The A2Zero goal is important."
    assert t["text_original"] == "The A20 goal is important."
    assert t["word_alignment_stale"] is True


def test_per_turn_speaker_beats_cluster_name(tmp_path):
    """How a reviewer says 'this one turn is somebody else' — the welded-turn case."""
    d = tmp_path / "proc"; d.mkdir()
    wb = _workbook(tmp_path, [_trow(TURNS[0], speaker="Missy Stults")],
                   [_srow("SPEAKER_00", "Brooks A Curtis", "roll call", confirm="y")])
    rep = parse_workbook(wb, TURNS)
    apply("v", rep, TURNS, SPEAKERS, wb, d)
    got = json.loads((d / "turns.reviewed.json").read_text())["turns"]
    assert got[0]["speaker_name"] == "Missy Stults" and got[0]["speaker_override"]
    assert got[2]["speaker_name"] == "Brooks A Curtis"   # same cluster, not overridden


def test_rejected_cluster_is_removed_from_speakers(tmp_path):
    d = tmp_path / "proc"; d.mkdir()
    wb = _workbook(tmp_path, [], [_srow("SPEAKER_00", "Brooks A Curtis", "roll call", confirm="n")])
    rep = parse_workbook(wb, TURNS)
    apply("v", rep, TURNS, SPEAKERS, wb, d)
    assert "SPEAKER_00" not in json.loads((d / "speakers.reviewed.json").read_text())


def test_confirmed_cluster_records_human_method(tmp_path):
    d = tmp_path / "proc"; d.mkdir()
    wb = _workbook(tmp_path, [], [_srow("SPEAKER_00", "Brooks A Curtis", "roll call",
                                        confirm="y", correct="Brooks Curtis")])
    rep = parse_workbook(wb, TURNS)
    apply("v", rep, TURNS, SPEAKERS, wb, d)
    s = json.loads((d / "speakers.reviewed.json").read_text())["SPEAKER_00"]
    assert s["method"] == "human" and s["confidence"] == 1.0 and s["review_confirmed"]
    assert "was 'Brooks A Curtis'" in s["evidence"]


# ───────────────────────── re-anchoring ─────────────────────────

def _corr(key, start_ms, end_ms):
    return {"key": key, "note": "x",
            "anchor": {"start_ms": start_ms, "end_ms": end_ms, "cluster": "SPEAKER_00"}}


def test_reanchor_survives_a_key_change():
    """S3 re-run with a new gap tolerance renumbers and re-bounds every turn. The
    correction must follow the audio, not the string."""
    new_turns = [{"key": "v:00000:9800-20400", "start": 9.8, "end": 20.4}]
    matched, lost = reanchor([_corr("v:00000:10000-20000", 10000, 20000)], new_turns)
    assert not lost
    assert matched[0]["reanchored"] and matched[0]["key_now"] == "v:00000:9800-20400"
    assert matched[0]["reanchor_iou"] > 0.9


def test_reanchor_refuses_a_turn_that_swallowed_its_neighbours():
    """A merged turn spanning 10s-120s overlaps the original but is a different scope.
    Attaching the correction there would silently widen what it applies to."""
    new_turns = [{"key": "v:00000:10000-120000", "start": 10.0, "end": 120.0}]
    matched, lost = reanchor([_corr("v:00000:10000-20000", 10000, 20000)], new_turns)
    assert not matched and lost and "overlap" in lost[0]["reason"]


def test_reanchor_prefers_exact_key_without_scanning():
    same = [{"key": "v:00001:21000-30000", "start": 21.0, "end": 30.0}]
    matched, _ = reanchor([_corr("v:00001:21000-30000", 21000, 30000)], same)
    assert matched[0]["reanchored"] is False


# ───────────────────────── person resolution + privacy ─────────────────────────

@pytest.fixture(scope="module")
def db():
    psycopg = pytest.importorskip("psycopg", reason="psycopg not installed")
    try:
        admin = psycopg.connect(f"host={HOST} port={PORT} user={USER} dbname=postgres",
                                autocommit=True, connect_timeout=5)
    except Exception as exc:                                       # pragma: no cover
        pytest.skip(f"no Postgres at {HOST}:{PORT} — {exc}")
    name = f"gv_test_{uuid.uuid4().hex[:10]}"
    admin.execute(f'CREATE DATABASE "{name}"')
    conn = psycopg.connect(f"host={HOST} port={PORT} user={USER} dbname={name}",
                           autocommit=True)
    conn.execute((SCHEMA_DIR / "claim_store.sql").read_text())
    conn.execute("""
        INSERT INTO persons (full_name, is_public_figure) VALUES
          ('Brooks Curtis', TRUE), ('Sara Nedrich', TRUE), ('Melissa Stults', TRUE),
          ('Judah Garber', TRUE), ('Jane Resident', FALSE), ('Chris Lee', TRUE),
          ('Chris Lee', TRUE)""")
    yield conn
    conn.close()
    admin.execute(f'DROP DATABASE "{name}" WITH (FORCE)')
    admin.close()


def test_exact_and_middle_initial_resolve_confidently(db):
    links = {l.cluster: l for l in resolve_people(
        db, {"a": "Brooks Curtis", "b": "Sara M Nedrich"})}
    assert links["a"].how == "exact" and links["a"].confident
    assert links["b"].how == "name_key" and links["b"].confident


def test_surname_only_never_auto_links(db):
    """'Ken Garber' is a public commenter with no persons row. Linking him to the
    unrelated official 'Judah Garber' would file a voiceprint under the wrong human."""
    l = resolve_people(db, {"a": "Ken Garber"})[0]
    assert not l.confident and l.person_id is None
    assert l.how == "unmatched"


def test_first_initial_match_is_proposed_not_linked(db):
    """Missy -> Melissa is almost certainly right, and 'almost certainly' is exactly
    the case that needs a human, not a threshold."""
    l = resolve_people(db, {"a": "Missy Stults"})[0]
    assert l.how == "proposed" and not l.confident and l.person_id is None
    assert [c["name"] for c in l.candidates] == ["Melissa Stults"]


def test_duplicate_names_are_ambiguous_not_arbitrary(db):
    l = resolve_people(db, {"a": "Chris Lee"})[0]
    assert l.how == "ambiguous" and l.person_id is None and len(l.candidates) == 2


def test_private_individual_is_flagged_not_public(db):
    l = resolve_people(db, {"a": "Jane Resident"})[0]
    assert l.confident and l.is_public_figure is False


def test_database_refuses_a_private_voiceprint(db):
    """Belt and braces: enroll_voiceprints checks is_public_figure, and the trigger
    enforces it. The trigger is the one that cannot be bypassed."""
    pid = db.execute("SELECT id FROM persons WHERE full_name='Jane Resident'").fetchone()[0]
    with pytest.raises(Exception, match="is_public_figure"):
        db.execute("INSERT INTO person_voiceprints (person_id, embedding, embedding_model,"
                   " enrolled_from) VALUES (%s,%s,'m','human')", (pid, str([0.1] * 256)))


def test_voiceprint_accepts_the_dimension_pyannote_actually_emits(db):
    """The schema said vector(192); community-1 emits 256. pgvector rejects a
    mismatch outright, so this would have failed at enrollment time."""
    pid = db.execute("SELECT id FROM persons WHERE full_name='Brooks Curtis'").fetchone()[0]
    db.execute("INSERT INTO person_voiceprints (person_id, embedding, embedding_model,"
               " enrolled_from) VALUES (%s,%s,'community-1','human')",
               (pid, str([0.01] * 256)))
    assert db.execute("SELECT vector_dims(embedding) FROM person_voiceprints "
                      "WHERE person_id=%s", (pid,)).fetchone()[0] == 256


def test_one_voiceprint_per_person_per_model(db):
    """Enrollment is incremental; without this constraint, re-processing a meeting
    yields two centroids for one person and every later lookup has to guess."""
    pid = db.execute("SELECT id FROM persons WHERE full_name='Sara Nedrich'").fetchone()[0]
    db.execute("INSERT INTO person_voiceprints (person_id, embedding, embedding_model,"
               " enrolled_from) VALUES (%s,%s,'community-1','human')",
               (pid, str([0.02] * 256)))
    with pytest.raises(Exception, match="person_voiceprints_person_model_key"):
        db.execute("INSERT INTO person_voiceprints (person_id, embedding, embedding_model,"
                   " enrolled_from) VALUES (%s,%s,'community-1','human')",
                   (pid, str([0.03] * 256)))


def test_alias_makes_a_proposed_match_resolve(db):
    from pipeline.import_review import approve_link
    approve_link(db, "Missy Stults", "Melissa Stults")
    l = resolve_people(db, {"a": "Missy Stults"})[0]
    assert l.how == "alias" and l.confident and l.is_public_figure


# ───────────────────────── enrollment gates ─────────────────────────

MODEL = "pyannote/speaker-diarization-community-1"


def _enroll(db, clusters, embeddings, video="vidA", commit=True):
    from pipeline.import_review import enroll_voiceprints
    corr = {"clusters": [{"cluster": c, "decision": d, "name": n}
                         for c, d, n in clusters]}
    links = resolve_people(db, {c: n for c, d, n in clusters if d == "confirmed" and n})
    return enroll_voiceprints(db, video, corr, links, "tester",
                              dry_run=not commit, embeddings=embeddings, model=MODEL)


def test_unconfirmed_cluster_is_never_enrolled(db):
    """The gate the Confirm column exists for."""
    out = _enroll(db, [("SPEAKER_A", "unreviewed", "Brooks Curtis")], {"SPEAKER_A": [0.1] * 256})
    assert out["enrolled"] == [] and out["refused"] == []


def test_private_individual_is_refused_before_the_trigger_fires(db):
    """The trigger is the backstop; this check exists so the refusal is legible in the
    report rather than arriving as an exception."""
    out = _enroll(db, [("SPEAKER_B", "confirmed", "Jane Resident")], {"SPEAKER_B": [0.1] * 256})
    assert out["enrolled"] == []
    assert "is_public_figure is FALSE" in out["refused"][0]["reason"]


def test_unresolvable_name_is_refused(db):
    out = _enroll(db, [("SPEAKER_C", "confirmed", "Ken Garber")], {"SPEAKER_C": [0.1] * 256})
    assert out["enrolled"] == []
    assert "not confidently resolved" in out["refused"][0]["reason"]


def test_missing_centroid_is_refused(db):
    out = _enroll(db, [("SPEAKER_D", "confirmed", "Brooks Curtis")], {})
    assert out["enrolled"] == []
    assert "no centroid" in out["refused"][0]["reason"]


def test_reenrolling_the_same_cluster_does_not_double_count(db):
    """Re-running the importer after correcting one more row must not fold the same
    audio in twice — that inflates sample_count, which is a claim about how many
    INDEPENDENT samples back the centroid."""
    who = [("SPEAKER_E", "confirmed", "Chris Solo")]
    db.execute("INSERT INTO persons (full_name, is_public_figure) VALUES ('Chris Solo', TRUE)")
    assert _enroll(db, who, {"SPEAKER_E": [0.1] * 256})["enrolled"][0]["action"] == "enrolled"
    again = _enroll(db, who, {"SPEAKER_E": [0.9] * 256})["enrolled"][0]
    assert again["action"] == "already_enrolled" and again["sample_count"] == 1


def test_a_second_meeting_does_increment(db):
    db.execute("INSERT INTO persons (full_name, is_public_figure) VALUES ('Dana Two', TRUE)")
    who = [("SPEAKER_F", "confirmed", "Dana Two")]
    _enroll(db, who, {"SPEAKER_F": [0.1] * 256}, video="vidA")
    out = _enroll(db, who, {"SPEAKER_F": [0.2] * 256}, video="vidB")["enrolled"][0]
    assert out["action"] == "updated" and out["sample_count"] == 2


def test_reattributed_cluster_is_reported_not_silently_moved(db):
    """If a cluster was enrolled as one person and is later confirmed as another, the
    first centroid contains a stranger's voice. Averaging is not reversible, so this
    has to surface rather than be patched over."""
    db.execute("INSERT INTO persons (full_name, is_public_figure) VALUES"
               " ('Early Guess', TRUE), ('Real Person', TRUE)")
    emb = {"SPEAKER_G": [0.3] * 256}
    _enroll(db, [("SPEAKER_G", "confirmed", "Early Guess")], emb, video="vidC")
    out = _enroll(db, [("SPEAKER_G", "confirmed", "Real Person")], emb, video="vidC")
    assert out["enrolled"] == []
    assert "previously enrolled as 'Early Guess'" in out["refused"][0]["reason"]


def test_dry_run_writes_nothing(db):
    db.execute("INSERT INTO persons (full_name, is_public_figure) VALUES ('Dry Run', TRUE)")
    before = db.execute("SELECT count(*) FROM person_voiceprints").fetchone()[0]
    out = _enroll(db, [("SPEAKER_H", "confirmed", "Dry Run")], {"SPEAKER_H": [0.4] * 256},
                  commit=False)
    assert out["dry_run"] and out["enrolled"][0]["action"] == "would enroll"
    assert db.execute("SELECT count(*) FROM person_voiceprints").fetchone()[0] == before


# ───────────────────────── export/import round trip ─────────────────────────

def test_refill_puts_every_correction_back_on_its_own_row(tmp_path, monkeypatch):
    """The end-to-end property that matters when the workbook format changes under a
    reviewer mid-pass: export -> annotate -> import -> RE-EXPORT -> every correction is
    on the same utterance it started on.

    Caught a real defect. refill() re-anchored against turns.json, which contains no
    interjection keys, so all six interjection corrections matched their PARENT turn by
    interval overlap and were written to the clerk's row instead of the answerer's.
    """
    import pipeline.export_review as ex
    from pipeline.import_review import parse_workbook, apply

    parent = {"seq": 0, "speaker": "SPEAKER_04", "start": 36.0, "end": 42.0,
              "text": "Commissioner Berkowitz? Commissioner Brown?",
              "key": "v:00000:36000-42000",
              "interjections": [{"speaker": "SPEAKER_07", "start": 38.8, "end": 40.0,
                                 "text": "Present in Ann Arbor."}]}
    plain = {"seq": 1, "speaker": "SPEAKER_05", "start": 45.0, "end": 50.0,
             "text": "Thank you.", "key": "v:00001:45000-50000", "interjections": []}
    turns = [parent, plain]
    speakers = {"SPEAKER_04": {"name": "Erin Donnelly", "method": "human",
                               "evidence": "x", "confidence": 1.0}}

    proc = tmp_path / "processing" / "vid"
    proc.mkdir(parents=True)
    (proc / "turns.json").write_text(json.dumps({"video_id": "vid", "turns": turns}))
    (proc / "speakers.json").write_text(json.dumps(speakers))
    monkeypatch.setattr(ex, "REPO", tmp_path)

    meta = {"title": "T", "date": "d", "duration_min": 1.0, "provenance": {}}
    stats = ex.speaker_stats(turns, speakers)
    wb_path = tmp_path / "rt.xlsx"
    ex.build_xlsx("vid", turns, speakers, stats, meta, wb_path)

    # Annotate the INTERJECTION row, which is the case that broke.
    from openpyxl import load_workbook
    wb = load_workbook(wb_path)
    ws = wb["Transcript"]
    int_row = next(r for r in range(2, ws.max_row + 1)
                   if "#int@" in str(ws.cell(row=r, column=ex.COL["utterance_key"]).value))
    ws.cell(row=int_row, column=ex.COL["✎ Speaker correction"], value="Mike Berkowitz")
    ws.cell(row=int_row, column=ex.COL["✎ Flag"], value="speaker")
    wb.save(wb_path)

    rep = parse_workbook(wb_path, turns)
    assert len(rep.turns) == 1
    apply("vid", rep, turns, speakers, wb_path, proc)

    # Re-export from scratch, then refill from corrections.json.
    fresh = tmp_path / "fresh.xlsx"
    ex.build_xlsx("vid", turns, speakers, stats, meta, fresh)
    out = ex.refill("vid", fresh, proc / "corrections.json")
    assert out["restored"] == 1 and out["lost"] == []
    assert out["reanchored"] == 0, "an unchanged turns.json must need no re-anchoring"

    ws2 = load_workbook(fresh, data_only=True)["Transcript"]
    landed = [r for r in range(2, ws2.max_row + 1)
              if ws2.cell(row=r, column=ex.COL["✎ Speaker correction"]).value]
    assert len(landed) == 1
    row = landed[0]
    assert "#int@" in str(ws2.cell(row=row, column=ex.COL["utterance_key"]).value), \
        "correction landed on the parent turn instead of the interjection"
    assert ws2.cell(row=row, column=ex.COL["Cluster"]).value == "SPEAKER_07"
